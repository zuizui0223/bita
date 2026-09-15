from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt

from scripts import build_bita_mechanism_candidate_sources as candidate
from scripts.format_ecology_submission_docx import (
    REF_BREAK,
    TITLE_BREAK,
    _format_document,
)


ROOT = Path(__file__).resolve().parents[1]


def _generated_parts() -> tuple[str, str, str]:
    text = candidate.build_main_source()
    title_page, body = text.split(candidate.TITLE_BREAK, 1)
    return text, title_page, body


def test_ecology_title_abstract_and_keyword_limits() -> None:
    text, title_page, body = _generated_parts()

    title = text.splitlines()[0].removeprefix("# ").strip()
    assert len(title) <= candidate.TITLE_CHAR_LIMIT

    abstract = candidate._abstract_text(body)
    assert candidate._word_count(abstract) <= candidate.ABSTRACT_WORD_LIMIT

    match = re.search(r"(?m)^\*\*Keywords:\*\*\s*(.+)$", title_page)
    assert match is not None
    keywords = [item.strip() for item in match.group(1).split(";")]
    assert candidate.KEYWORD_MIN <= len(keywords) <= candidate.KEYWORD_MAX
    assert keywords == sorted(keywords, key=str.casefold)
    assert len({item.casefold() for item in keywords}) == len(keywords)
    assert keywords == [
        "antagonism",
        "causal identification",
        "ecological mechanism",
        "factorial experiment",
        "floral defence",
        "partial identification",
        "pollination",
        "trait interaction",
    ]


def test_ecology_required_title_page_items_precede_abstract() -> None:
    text, title_page, body = _generated_parts()

    assert "**Journal:** Ecology" in title_page
    assert "**Manuscript type:** Concepts & Synthesis" in title_page
    assert "**Authors and affiliations:** [Author-controlled]" in title_page
    assert candidate.CORRESPONDING_AUTHOR_LINE in title_page
    assert candidate.OPEN_RESEARCH_STATEMENT in title_page
    assert "**Keywords:**" in title_page

    assert "## Abstract" not in title_page
    assert body.lstrip().startswith("## Abstract")
    assert "**Keywords:**" not in body
    assert "**Open Research statement:**" not in body
    assert text.count("**Keywords:**") == 1
    assert text.count("**Open Research statement:**") == 1


def test_ecology_formatter_contract() -> None:
    doc = Document()
    doc.add_paragraph("Title page")
    doc.add_paragraph(TITLE_BREAK)
    doc.add_paragraph("Abstract text")
    doc.add_paragraph(REF_BREAK)
    doc.add_paragraph("Figure material")

    _format_document(doc, appendix=False)

    for section in doc.sections:
        assert section.page_width == Inches(8.5)
        assert section.page_height == Inches(11)
        assert section.top_margin == Inches(1)
        assert section.bottom_margin == Inches(1)
        assert section.left_margin == Inches(1)
        assert section.right_margin == Inches(1)

    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == 2.0

    xml = doc._element.xml
    assert "w:lnNumType" in xml
    assert 'w:restart="continuous"' in xml
    assert " PAGE " in doc.sections[0].footer.paragraphs[0]._p.xml


def test_ecology_package_and_ai_disclosure_are_tracked() -> None:
    plan = (ROOT / "submission" / "ECOLOGY_UPLOAD_PACKAGE_PLAN.md").read_text(encoding="utf-8")
    audit = (ROOT / "submission" / "ECOLOGY_CONCEPTS_SYNTHESIS_FIT_AUDIT.md").read_text(encoding="utf-8")
    portal = (ROOT / "submission" / "AUTHOR_AND_PORTAL_METADATA_TEMPLATE.md").read_text(encoding="utf-8")

    for text in (plan, audit):
        assert "April 2026" in text
        assert "21" in text and "10" in text
        assert "PACKAGE_QA_RECEIPT.txt" in text
        assert "Open Research" in text

    assert "Acknowledgments" in portal
    assert "submission form" in portal
    assert "AI" in portal


def test_final_submission_documents_reflect_compliance_closure() -> None:
    cover = (ROOT / "submission" / "COVER_LETTER_ECOLOGY_CONCEPTS_SYNTHESIS.md").read_text(encoding="utf-8")
    checklist = (ROOT / "submission" / "SUBMISSION_CHECKLIST.md").read_text(encoding="utf-8")
    audit = (ROOT / "docs" / "FINAL_SUBMISSION_AUDIT.md").read_text(encoding="utf-8")

    assert "will be rebuilt" not in cover.casefold()
    assert "30-page Main / 38-page Appendix" not in cover
    assert "**21 pages**" in cover
    assert "standard **30-page**" in cover
    assert "no over-length justification" in cover

    assert "ECOLOGY_APR2026_AUTOMATED_COMPLIANCE = PASS" in checklist
    assert "ecology_apr2026_title_page=PASS" in checklist
    assert "AI_DISCLOSURE = BLOCKED_AUTHOR_APPROVAL_IF_REQUIRED" in checklist
    assert "live author instructions / upload portal" in checklist

    assert "ECOLOGY_APR2026_AUTOMATED_COMPLIANCE = PASS" in audit
    assert "ecology_apr2026_title_page=PASS" in audit
    assert "READY_FOR_AUTHOR_METADATA" in audit
    assert "SUBMITTED" in audit
